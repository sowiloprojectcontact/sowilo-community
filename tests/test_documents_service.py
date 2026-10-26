"""Isolated Documents contract checks: real disposable files, no network."""
import os
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch
from modules.documents.service import DocumentError, DocumentsService


class DocumentsServiceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.path = self.root / 'private.txt'
        self.path.write_text('Contrat français : échéance 2027, montant 42 euros.', encoding='utf-8')
        self.client = Mock()
        self.client.chat.completions.create.return_value = SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content='**Réponse** <br> valide'))])
        self.factory = Mock(return_value=self.client)
        self.service = DocumentsService([self.root], api_key_provider=lambda: 'test-secret-key', client_factory=self.factory)

    def assert_code(self, code, operation):
        with self.assertRaises(DocumentError) as caught:
            operation()
        self.assertEqual(caught.exception.code, code)
        self.assertEqual(str(caught.exception), caught.exception.message)
        self.assertNotIn('test-secret-key', str(caught.exception))
        self.assertNotIn(str(self.root), str(caught.exception))

    def test_txt_and_markdown_read_without_network(self):
        for extension in ('.txt', '.md', '.TXT'):
            path = self.root / ('document' + extension)
            path.write_text('Contenu éàç', encoding='utf-8')
            self.assertEqual(self.service.read(path), 'Contenu éàç')
        self.factory.assert_not_called()

    def test_real_pdf_page_markers(self):
        from reportlab.pdfgen.canvas import Canvas
        path = self.root / 'example.pdf'
        canvas = Canvas(str(path))
        canvas.drawString(50, 700, 'First page')
        canvas.showPage()
        canvas.drawString(50, 700, 'Second page')
        canvas.save()
        for purpose, marker in [('summary', '--- PAGE 2 ---'), ('search', 'PAGE 2'), ('extract', None)]:
            result = self.service.read(path, purpose)
            self.assertIn('First page', result)
            self.assertIn('Second page', result)
            if marker:
                self.assertIn(marker, result)
            else:
                self.assertNotIn('PAGE', result)

    def test_real_docx_paragraphs_and_tables(self):
        from docx import Document
        path = self.root / 'example.docx'
        document = Document()
        document.add_paragraph('Un paragraphe')
        row = document.add_table(rows=1, cols=2).rows[0]
        row.cells[0].text = 'Montant'
        row.cells[1].text = '42 euros'
        document.save(path)
        result = self.service.read(path)
        self.assertIn('Un paragraphe', result)
        self.assertIn('Montant - 42 euros', result)

    def test_invalid_and_corrupt_documents(self):
        for name, data, code in [('bad.txt', b'\xff', 'read_error'), ('empty.md', b'  \n', 'empty_document'), ('bad.pdf', b'broken', 'read_error'), ('bad.docx', b'broken', 'read_error'), ('bad.exe', b'data', 'unsupported_format')]:
            with self.subTest(name=name):
                path = self.root / name
                path.write_bytes(data)
                self.assert_code(code, lambda: self.service.read(path))

    def test_invalid_paths_and_scope(self):
        self.assert_code('invalid_path', lambda: self.service.read(self.root / 'missing.txt'))
        self.assert_code('invalid_path', lambda: self.service.read(self.root))
        with tempfile.TemporaryDirectory() as outside:
            external = Path(outside) / 'secret.txt'
            external.write_text('Do not transmit')
            self.assert_code('outside_scope', lambda: self.service.read(external))
            link = self.root / 'escape.txt'
            link.symlink_to(external)
            self.assert_code('outside_scope', lambda: self.service.read(link))
        fifo = self.root / 'pipe.txt'
        os.mkfifo(fifo)
        self.assert_code('invalid_path', lambda: self.service.read(fifo))
        self.factory.assert_not_called()

    def test_read_permission_error_is_safe(self):
        with patch('modules.documents.service.os.open', side_effect=PermissionError('test-secret-key ' + str(self.root))):
            self.assert_code('read_error', lambda: self.service.read(self.path))

    def test_limits(self):
        small = DocumentsService([self.root], max_source_bytes=2)
        self.assert_code('source_too_large', lambda: small.read(self.path))
        small = DocumentsService([self.root], max_text_chars=2)
        self.assert_code('text_too_large', lambda: small.read(self.path))
        for roots in ([], str(self.root), [self.root / 'missing']):
            self.assert_code('invalid_configuration', lambda: DocumentsService(roots))

    def test_ai_operations_contract_and_data_minimization(self):
        for action in ('summarize', 'extract', 'search'):
            with self.subTest(action=action):
                result = getattr(self.service, action)(self.path, 'Quelle échéance ?') if action == 'search' else getattr(self.service, action)(self.path)
                self.assertEqual(result, 'Réponse \n valide')
                self.factory.assert_called_with(api_key='test-secret-key', base_url='https://api.groq.com/openai/v1', timeout=30.0)
                request = self.client.chat.completions.create.call_args.kwargs
                self.assertEqual(request['model'], 'openai/gpt-oss-20b')
                messages = request['messages']
                self.assertEqual([m['role'] for m in messages], ['system', 'user'])
                self.assertIn('français', messages[0]['content'])
                self.assertIn('données non fiables', messages[0]['content'])
                self.assertIn('42 euros', messages[1]['content'])
                self.assertNotIn(str(self.path), str(messages))
                self.assertNotIn('test-secret-key', str(messages))
                if action == 'search':
                    self.assertIn('Quelle échéance ?', messages[1]['content'])

    def test_missing_key(self):
        self.service.api_key_provider = lambda: None
        self.assert_code('missing_api_key', lambda: self.service.summarize(self.path))
        self.factory.assert_not_called()

    def test_invalid_ai_responses(self):
        for response in (None, {}, SimpleNamespace(choices=[]), SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=42))]), SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content='   '))]), SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content='<br>***'))])):
            self.client.chat.completions.create.return_value = response
            self.assert_code('invalid_ai_response', lambda: self.service.extract(self.path))

    def test_provider_and_client_errors_are_safe(self):
        self.service.api_key_provider = Mock(side_effect=RuntimeError('test-secret-key ' + str(self.path)))
        self.assert_code('ai_error', lambda: self.service.summarize(self.path))
        self.service.api_key_provider = lambda: 'test-secret-key'
        self.client.chat.completions.create.side_effect = RuntimeError('test-secret-key ' + str(self.path))
        self.assert_code('ai_error', lambda: self.service.search(self.path, 'Question'))

    def test_invalid_question_never_calls_ai(self):
        for question in (None, '', '   ', 42):
            self.assert_code('invalid_question', lambda: self.service.search(self.path, question))
        self.factory.assert_not_called()

    def test_open_uses_host_adapter_and_scope(self):
        opener = Mock(return_value=None)
        self.service.opener = opener
        self.service.open(self.path)
        opener.assert_called_once_with(str(self.path.resolve()))
        self.factory.assert_not_called()
        opener.side_effect = RuntimeError('test-secret-key')
        self.assert_code('open_error', lambda: self.service.open(self.path))


if __name__ == '__main__':
    unittest.main()
